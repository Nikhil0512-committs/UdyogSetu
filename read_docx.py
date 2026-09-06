import docx
import sys

def read_docx(file_path):
    try:
        # Reconfigure stdout to use utf-8 to avoid charmap errors on Windows
        sys.stdout.reconfigure(encoding='utf-8')
        doc = docx.Document(file_path)
        
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
                
        for table in doc.tables:
            print("\n--- TABLE ---")
            for row in table.rows:
                row_data = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                print(" | ".join(row_data))
            print("-------------\n")
            
    except Exception as e:
        print(f"Error reading {file_path}: {e}")

if __name__ == '__main__':
    read_docx(sys.argv[1])
